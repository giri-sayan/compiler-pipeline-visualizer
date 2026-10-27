import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def style_table_header(row, col_widths, bg_hex="0F172A"):
    for idx, cell in enumerate(row.cells):
        set_cell_background(cell, bg_hex)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        cell.width = col_widths[idx]
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9.5)

def style_table_data(row, col_widths, is_even=False):
    bg_hex = "F8FAFC" if is_even else "FFFFFF"
    for idx, cell in enumerate(row.cells):
        set_cell_background(cell, bg_hex)
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        cell.width = col_widths[idx]
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(15, 23, 42)

def build_docx():
    doc = Document()
    img_dir = "C:/Users/user/.gemini/antigravity/brain/641c8502-4f1b-4291-8ee4-3619007b89ed"

    # Configure Margins: 1 inch everywhere
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Palette
    COLOR_PRIMARY = RGBColor(14, 165, 233)   # #0EA5E9 (Cyan/Teal)
    COLOR_NAVY = RGBColor(15, 23, 42)        # #0F172A
    COLOR_SLATE = RGBColor(71, 85, 105)      # #475569
    COLOR_GREEN = RGBColor(16, 185, 129)     # #10B981

    # ==================== COVER / HEADER BANNER ====================
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_t = title_p.add_run("Compiler Pipeline Visualizer (CPV)")
    run_t.font.name = "Arial"
    run_t.font.size = Pt(26)
    run_t.font.bold = True
    run_t.font.color.rgb = COLOR_NAVY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("An Interactive Multi-Language Compiler Engine & Execution Sandbox")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(14)
    run_sub.font.bold = True
    run_sub.font.color.rgb = COLOR_PRIMARY

    # Metadata Box Table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_widths = [Inches(2.2), Inches(4.3)]
    
    meta_data = [
        ("Course & Code:", "BCSE307P — Compiler Design"),
        ("Evaluation Phase:", "Review 3 (Final Comprehensive Review)"),
        ("Project Team:", "Ansuman Singh (24BCE2013) & Sayan Giri (24BCT0281)"),
        ("Live Production URL:", "https://compiler-pipeline-visualizer-mocha.vercel.app/"),
        ("GitHub Repository:", "https://github.com/giri-sayan/compiler-pipeline-visualizer")
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = meta_widths[0], meta_widths[1]
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "F8FAFC")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        
        p0 = c0.paragraphs[0]
        p0.add_run(label).font.bold = True
        p0.runs[0].font.size = Pt(9.5)
        p0.runs[0].font.color.rgb = COLOR_NAVY
        
        p1 = c1.paragraphs[0]
        p1.add_run(val).font.size = Pt(9.5)
        p1.runs[0].font.color.rgb = COLOR_NAVY
        if "http" in val:
            p1.runs[0].font.color.rgb = COLOR_PRIMARY
            p1.runs[0].font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    def add_section_header(text, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.bold = True
        if level == 1:
            run.font.size = Pt(16)
            run.font.color.rgb = COLOR_NAVY
            # Add a bottom border or underline effect with spacing
        elif level == 2:
            run.font.size = Pt(13)
            run.font.color.rgb = COLOR_PRIMARY
        else:
            run.font.size = Pt(11)
            run.font.color.rgb = COLOR_NAVY
        return p

    def add_body_p(text, bold_prefix=None, space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_b = p.add_run(bold_prefix)
            r_b.font.name = "Arial"
            r_b.font.bold = True
            r_b.font.size = Pt(10)
            r_b.font.color.rgb = COLOR_NAVY
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_NAVY
        return p

    def add_bullet_p(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r_b = p.add_run(bold_prefix)
        r_b.font.name = "Arial"
        r_b.font.bold = True
        r_b.font.size = Pt(10)
        r_b.font.color.rgb = COLOR_NAVY
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_NAVY
        return p

    def add_image_with_caption(img_filename, caption_text, width_inches=6.0):
        full_path = os.path.join(img_dir, img_filename)
        if os.path.exists(full_path):
            img_p = doc.add_paragraph()
            img_p.paragraph_format.space_before = Pt(10)
            img_p.paragraph_format.space_after = Pt(4)
            img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            img_p.paragraph_format.keep_with_next = True
            run_img = img_p.add_run()
            run_img.add_picture(full_path, width=Inches(width_inches))

            cap_p = doc.add_paragraph()
            cap_p.paragraph_format.space_before = Pt(2)
            cap_p.paragraph_format.space_after = Pt(12)
            cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_cap = cap_p.add_run(f"Figure: {caption_text}")
            run_cap.font.name = "Arial"
            run_cap.font.size = Pt(9)
            run_cap.font.italic = True
            run_cap.font.color.rgb = COLOR_SLATE

    # ==================== 1. EXECUTIVE SUMMARY ====================
    add_section_header("1. Executive Summary")
    add_body_p(
        "The Compiler Pipeline Visualizer (CPV) is an interactive, multi-language compiler engineering platform and educational sandbox designed to demystify every phase of program compilation and execution. While traditional computer science curricula treat compilers as opaque 'black boxes' (e.g. gcc file.c -o file), this project makes the entire compilation lifecycle visible, inspectable, and interactive."
    )
    add_body_p(
        "Across three developmental reviews, the system evolved from a small prototype into an industrial-strength, cloud-deployed educational IDE capable of processing 5,000-line programs across six programming languages with live virtual machine execution:"
    )

    # Evolution Table
    evo_tbl = doc.add_table(rows=8, cols=4)
    evo_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    evo_widths = [Inches(1.5), Inches(1.4), Inches(1.7), Inches(1.9)]
    
    evo_headers = ["Capability", "Review 1 Baseline", "Review 2 Milestone", "Review 3 (Final Deliverable)"]
    style_table_header(evo_tbl.rows[0], evo_widths)
    for c_idx, h in enumerate(evo_headers):
        evo_tbl.rows[0].cells[c_idx].paragraphs[0].text = h
    style_table_header(evo_tbl.rows[0], evo_widths)

    evo_rows = [
        ("Language Coverage", "Toy language only", "C subset (types, main, printf)", "6 Languages: C, C++, Java, Python, JavaScript, Toy Lang"),
        ("Input Scalability", "~20 lines", "~50 lines", "5,000+ lines processed & executed in <200 ms"),
        ("Program Execution", "None (Static)", "None (Static code generation only)", "Live In-Memory TAC Virtual Machine & Python Subprocess"),
        ("File Ingestion", "Manual typing", "Basic textarea editing", "Direct File Upload (.c, .cpp, .java, etc.) & Lossless Paste"),
        ("Theme & Contrast", "Single Light Theme", "Experimental Dark Mode (Low contrast)", "Dual Themes (Light & Night) with WCAG AA compliance (>=4.5:1)"),
        ("Line Synchronization", "Unchecked", "Frozen at lines 1..27 on large files", "Absolute Viewport Gutter Tracking up to line 5,000+"),
        ("Automated Testing", "8 unit tests", "17 unit tests", "136 passing regression & stress tests (100% pass rate)")
    ]
    for r_idx, r_data in enumerate(evo_rows, start=1):
        row = evo_tbl.rows[r_idx]
        for c_idx, val in enumerate(r_data):
            row.cells[c_idx].paragraphs[0].text = val
        style_table_data(row, evo_widths, is_even=(r_idx % 2 == 1))

    # ==================== 2. SYSTEM ARCHITECTURE ====================
    add_section_header("2. System Architecture & Polyglot Pipeline")
    add_body_p(
        "CPV is built around a pluggable, modular compiler frontend engine that normalizes source code from distinct programming paradigms into a unified intermediate analysis and code generation pipeline."
    )
    add_bullet_p("1. Ingestion Layer: ", "Accepts source code through direct file upload (.c, .cpp, .java, .js, .py, .toy, .txt), a lossless clipboard paste interceptor, or 41 isolated language presets.")
    add_bullet_p("2. Pluggable Frontend Layer: ", "Dispatches code through get_compiler(lang) into dedicated language frontends for C (C99), C++, Java, Python (ast.parse), JavaScript (ES6), and Toy Lang.")
    add_bullet_p("3. Unified Compiler Pipeline: ", "Converts tokens into a recursive AST, validates declarations in the Symbol Table, generates Three-Address Code (TAC) quadruples, and emits x86-64 machine code and RISC assembly.")
    add_bullet_p("4. Dual Execution Engines: ", "Executes generated TAC quadruples in a pure-Python in-memory Virtual Machine (tac_vm.py) with a 20,000-step instruction budget, plus an isolated native Python subprocess runner (executor.py).")

    add_image_with_caption("screenshot_pipeline_light.png", "Live Cloud Production Interface on Vercel displaying the Multi-Stage Compiler Pipeline")

    # ==================== 3. STAGE-BY-STAGE IMPLEMENTATION ====================
    add_section_header("3. Stage-by-Stage Implementation & Visual Verification")

    add_section_header("Stage 1: Lexical Analysis (Scanner)", level=2)
    add_body_p(
        "The lexer converts raw character streams into classified token objects (PREPROCESSOR, TYPE, IDENTIFIER, NUMBER, STRING, OPERATOR, DELIMITER). To guarantee zero UI lag when processing 5,000 lines, tokens are grouped into line-indexed dictionaries in O(N) linear time."
    )
    add_image_with_caption("screenshot_scanner_light.png", "Stage 1 Lexical Scanner: Classified Token Stream with Token Breakdown Pills")

    add_section_header("Stage 2: Syntax Analysis & Recursive-Descent Parsing", level=2)
    add_body_p(
        "Constructs a hierarchical Abstract Syntax Tree (AST) enforcing operator precedence. The C parser handles function prototypes ending in semicolons, comma-separated declarations, type qualifiers (const, unsigned, long), pointer declarators, and arrays. When syntax errors occur, the parser flags the exact line and column, rendering a diagnostic chip with Jump to Error navigation."
    )
    add_image_with_caption("screenshot_syntax_error.png", "Stage 2 Syntax Analysis: Parser Diagnostic Error Chip with Jump to Error Navigation")

    add_section_header("Stage 3: Semantic Analysis & Symbol Table", level=2)
    add_body_p(
        "Constructs a multi-level scoped Symbol Table tracking global vs function scope, enforcing declaration-before-use rules with typo suggestions (e.g. suggesting printf for rintf), detecting uninitialized variables, and validating standard C library function calls."
    )
    add_image_with_caption("screenshot_semantic_table.png", "Stage 3 Semantic Analysis: Scoped Symbol Table with Scope Lifetimes and Status Metadata")

    add_section_header("Stage 4 & 5: Intermediate Code (TAC) & Target Code Generation", level=2)
    add_body_p(
        "Linearizes AST nodes into canonical Three-Address Code quadruples (op, arg1, arg2, result). These quadruples are then translated into two target architectures: realistic Intel x86-64 assembly using System V AMD64 ABI calling conventions, and educational 3-register RISC pseudo-assembly (LOAD, ADD, STORE)."
    )

    add_section_header("Stage 6: In-Memory TAC Virtual Machine Execution", level=2)
    add_body_p(
        "To allow students to observe program execution without installing external compiler toolchains (such as gcc, javac, or node), CPV includes an in-memory TAC Virtual Machine (tac_vm.py). The VM simulates CPU registers, call frames, and memory with a 20,000-step instruction budget to terminate infinite loops safely."
    )

    # ==================== 4. UI INNOVATIONS & THEMES ====================
    add_section_header("4. UI/UX Innovations, Night Theme & Stage Diff Mode")
    add_body_p(
        "Review 3 introduced complete WCAG AA compliant Night Theme (Dark Mode) design tokens. Pale AST node backgrounds that previously caused text unreadability were re-engineered with high-contrast color tokens (>=4.5:1 contrast ratio)."
    )
    add_image_with_caption("screenshot_ast_dark.png", "AST Visual Explorer in Night Theme (Dark Mode) with WCAG AA Compliant Nodes")

    add_body_p(
        "Stage Diff Mode provides side-by-side comparative analysis between two programs (e.g. comparing unoptimized vs optimized code, or iterative vs recursive implementations), highlighting additions, deletions, and structural changes across AST and TAC quadruples."
    )
    add_image_with_caption("screenshot_diff_dark.png", "Multi-Language Stage Diff Mode: Side-by-Side AST and TAC Structural Comparison")

    # ==================== 5. CRITICAL BUG FIXES ====================
    add_section_header("5. Critical Engineering Bug Fixes: 5,000-Line Input Processing")
    add_bullet_p("1. Clipboard Buffer Truncation: ", "Windows console pipes clip clipboard copies to 4,096 bytes (~96 lines). Pasting large files prematurely truncated source code, causing syntax errors at line 96. Resolved by implementing a dedicated clipboard paste interceptor (reading e.clipboardData directly) and a native 'Upload Source File' feature supporting .c, .cpp, .java, .js, .py, .toy, .txt.")
    add_bullet_p("2. Flexbox 100,727px Gutter Clamping: ", "In large files, the line-number gutter expanded to 100,727 pixels inside an unconstrained flex container. Because clientHeight equaled scrollHeight, maxScrollTop was 0, clamping gutter.scrollTop to 0 and freezing visible line numbers at 1..27. Resolved by absolute viewport tracking (position: absolute; top: 0; bottom: 0; left: 0), dynamically adjusting gutter width, and synchronizing Jump-to-Error navigation.")

    # ==================== 6. SCALABILITY BENCHMARKS ====================
    add_section_header("6. Performance Benchmarks across 6 Languages")
    add_body_p(
        "Live measurements on valid 5,000-line programs across all six languages demonstrate sub-200ms compilation and execution times:"
    )

    bench_tbl = doc.add_table(rows=7, cols=5)
    bench_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    bench_widths = [Inches(1.5), Inches(1.1), Inches(1.6), Inches(1.2), Inches(1.1)]

    b_headers = ["Language", "Total Lines", "Full Pipeline Status", "Output", "Measured Time"]
    style_table_header(bench_tbl.rows[0], bench_widths)
    for c_idx, h in enumerate(b_headers):
        bench_tbl.rows[0].cells[c_idx].paragraphs[0].text = h
    style_table_header(bench_tbl.rows[0], bench_widths)

    bench_rows = [
        ("C (C99 Subset)", "5,002 lines", "Passed (All 6 Stages)", "Total: 3", "0.089 s (89 ms)"),
        ("C++", "5,002 lines", "Passed (All 6 Stages)", "Result: 3", "0.073 s (73 ms)"),
        ("Java", "5,001 lines", "Passed (All 6 Stages)", "Java: 3", "0.074 s (74 ms)"),
        ("Python", "5,000 lines", "Passed (All 6 Stages)", "Python: 3", "0.190 s (190 ms)"),
        ("JavaScript", "5,001 lines", "Passed (All 6 Stages)", "JS: 3", "0.070 s (70 ms)"),
        ("Toy Lang", "5,000 lines", "Passed (All 6 Stages)", "3", "0.159 s (159 ms)")
    ]
    for r_idx, r_data in enumerate(bench_rows, start=1):
        row = bench_tbl.rows[r_idx]
        for c_idx, val in enumerate(r_data):
            row.cells[c_idx].paragraphs[0].text = val
        style_table_data(row, bench_widths, is_even=(r_idx % 2 == 1))

    # ==================== 7. TEST SUITE ====================
    add_section_header("7. Comprehensive Automated Test Suite (136 Tests)")
    add_body_p(
        "The project is backed by 136 automated regression and stress tests achieving a 100% pass rate in 13.30 seconds:"
    )

    test_tbl = doc.add_table(rows=7, cols=4)
    test_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_widths = [Inches(2.2), Inches(0.8), Inches(2.6), Inches(0.9)]

    t_headers = ["Test Suite File", "Tests", "Scope & Verifications", "Status"]
    style_table_header(test_tbl.rows[0], t_widths)
    for c_idx, h in enumerate(t_headers):
        test_tbl.rows[0].cells[c_idx].paragraphs[0].text = h
    style_table_header(test_tbl.rows[0], t_widths)

    t_rows = [
        ("test_pipeline.py", "43", "Core 6-stage pipeline, Lexer, Parser, AST, TAC, ASM", "PASSED"),
        ("test_review3_upgrades.py", "42", "6 languages, TAC VM, Night theme, WCAG AA contrast", "PASSED"),
        ("test_final_review_stress.py", "28", "3k & 4k line benchmarks, high-line diagnostics", "PASSED"),
        ("test_diff_mode.py", "11", "Side-by-side AST & TAC diff visual comparisons", "PASSED"),
        ("test_paste_and_large_input_fix.py", "8", "5,000-line C stress test, parser prototypes & types", "PASSED"),
        ("test_gutter_line_numbers.py", "4", "Browser DOM viewport gutter tracking (Edge CDP)", "PASSED")
    ]
    for r_idx, r_data in enumerate(t_rows, start=1):
        row = test_tbl.rows[r_idx]
        for c_idx, val in enumerate(r_data):
            row.cells[c_idx].paragraphs[0].text = val
        style_table_data(row, t_widths, is_even=(r_idx % 2 == 1))

    # ==================== 8. CLOUD DEPLOYMENT ====================
    add_section_header("8. Cloud Production Deployment on Vercel")
    add_body_p(
        "The Compiler Pipeline Visualizer is deployed live in production on Vercel as a Python Serverless Function (vercel.json, api/index.py). Every push to the main branch on GitHub automatically triggers build verification and zero-downtime deployment."
    )
    add_bullet_p("Production URL: ", "https://compiler-pipeline-visualizer-mocha.vercel.app/")
    add_bullet_p("GitHub Repository: ", "https://github.com/giri-sayan/compiler-pipeline-visualizer")

    # ==================== 9. CONCLUSION ====================
    add_section_header("9. Conclusion & Final Deliverables")
    add_body_p(
        "All faculty expectations and Review 3 objectives have been successfully fulfilled: 6 supported programming languages, 5,000+ line input scalability, in-memory virtual machine execution, WCAG AA compliant Night Theme, synchronized viewport gutter, and 136 passing automated tests."
    )

    output_path = "Compiler_Pipeline_Visualizer_Project_Report.docx"
    doc.save(output_path)
    print(f"Report saved successfully as {output_path} ({os.path.getsize(output_path)} bytes)")

if __name__ == "__main__":
    build_docx()
