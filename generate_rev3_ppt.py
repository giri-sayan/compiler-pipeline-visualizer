import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    BG_LIGHT = RGBColor(248, 250, 252)        # #F8FAFC
    CARD_BG = RGBColor(255, 255, 255)         # White
    CARD_BG_ALT = RGBColor(241, 245, 249)     # #F1F5F9
    CARD_BORDER = RGBColor(226, 232, 240)     # #E2E8F0
    TEXT_NAVY = RGBColor(15, 23, 42)          # #0F172A
    TEXT_MUTED = RGBColor(100, 116, 139)      # #64748B
    PRIMARY_BLUE = RGBColor(14, 165, 233)     # #0EA5E9 (Cyan/Blue)
    ACCENT_TEAL = RGBColor(13, 148, 136)      # #0D9488
    DARK_NAVY = RGBColor(15, 23, 42)          # Dark background
    SUCCESS_GREEN = RGBColor(16, 185, 129)    # #10B981
    WARNING_AMBER = RGBColor(245, 158, 11)    # #F59E0B
    PURPLE_ACCENT = RGBColor(99, 102, 241)    # #6366F1

    def set_slide_background(slide, color):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, title_text, category_text=None):
        # Category / breadcrumb pill
        if category_text:
            cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.35))
            tf_cat = cat_box.text_frame
            tf_cat.word_wrap = True
            tf_cat.margin_left = tf_cat.margin_right = tf_cat.margin_top = tf_cat.margin_bottom = 0
            p_cat = tf_cat.paragraphs[0]
            p_cat.text = category_text.upper()
            p_cat.font.size = Pt(10)
            p_cat.font.bold = True
            p_cat.font.color.rgb = PRIMARY_BLUE

        # Main Slide Title
        top_pos = Inches(0.7) if category_text else Inches(0.5)
        title_box = slide.shapes.add_textbox(Inches(0.8), top_pos, Inches(11.5), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = TEXT_NAVY

        # Accent Bar under title
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.8), top_pos + Inches(0.75), Inches(1.2), Inches(0.06)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = PRIMARY_BLUE
        bar.line.color.rgb = PRIMARY_BLUE

    # ==========================================
    # SLIDE 1: Title Slide (Dark / Executive Tech)
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, DARK_NAVY)

    # Accent decorative glow rectangle
    accent_strip = slide1.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.8), Inches(1.5), Inches(0.12), Inches(4.5)
    )
    accent_strip.fill.solid()
    accent_strip.fill.fore_color.rgb = PRIMARY_BLUE
    accent_strip.line.fill.background()

    # Title Box
    t_box = slide1.shapes.add_textbox(Inches(1.2), Inches(1.5), Inches(11), Inches(2.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "Compiler Pipeline Visualizer"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)

    p1_sub = tf1.add_paragraph()
    p1_sub.text = "An Interactive Multi-Language Compiler Engine & Execution Sandbox"
    p1_sub.font.size = Pt(22)
    p1_sub.font.bold = True
    p1_sub.font.color.rgb = PRIMARY_BLUE
    p1_sub.space_before = Pt(12)

    # Review 3 Badge
    badge_box = slide1.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(1.2), Inches(4.0), Inches(5.8), Inches(0.55)
    )
    badge_box.fill.solid()
    badge_box.fill.fore_color.rgb = RGBColor(30, 41, 59)
    badge_box.line.color.rgb = PRIMARY_BLUE
    b_tf = badge_box.text_frame
    b_p = b_tf.paragraphs[0]
    b_p.text = "Review 3 (Final Review)  |  Compiler Design  |  BCSE307P"
    b_p.font.size = Pt(13)
    b_p.font.bold = True
    b_p.font.color.rgb = RGBColor(241, 245, 249)
    b_p.alignment = PP_ALIGN.CENTER

    # Project Team
    author_box = slide1.shapes.add_textbox(Inches(1.2), Inches(4.8), Inches(6.0), Inches(1.8))
    atf = author_box.text_frame
    ap0 = atf.paragraphs[0]
    ap0.text = "PRESENTED BY:"
    ap0.font.size = Pt(11)
    ap0.font.bold = True
    ap0.font.color.rgb = TEXT_MUTED

    ap1 = atf.add_paragraph()
    ap1.text = "Ansuman Singh  (24BCE2013)"
    ap1.font.size = Pt(16)
    ap1.font.bold = True
    ap1.font.color.rgb = RGBColor(255, 255, 255)

    ap2 = atf.add_paragraph()
    ap2.text = "Sayan Giri  (24BCT0281)"
    ap2.font.size = Pt(16)
    ap2.font.bold = True
    ap2.font.color.rgb = RGBColor(255, 255, 255)

    # Key Highlights Card on Right
    hl_card = slide1.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(7.5), Inches(2.2), Inches(5.0), Inches(4.2)
    )
    hl_card.fill.solid()
    hl_card.fill.fore_color.rgb = RGBColor(30, 41, 59)
    hl_card.line.color.rgb = RGBColor(51, 65, 85)
    htf = hl_card.text_frame
    htf.word_wrap = True
    htf.margin_left = htf.margin_right = Inches(0.4)
    htf.margin_top = Inches(0.35)

    hp0 = htf.paragraphs[0]
    hp0.text = "REVIEW 3 CORE DELIVERABLES"
    hp0.font.size = Pt(13)
    hp0.font.bold = True
    hp0.font.color.rgb = PRIMARY_BLUE

    highlights = [
        "1. Large-Input Scaling: 5,000+ source lines compiled & executed in <200ms",
        "2. Multi-Language Frontends: C (C99), C++, Java, Python, JS, Toy Lang",
        "3. In-Memory Execution: Sandboxed TAC Virtual Machine & Python Runner",
        "4. Source Workspace: Native File Upload (.c, .cpp, .java, .js, .py, .toy, .txt)",
        "5. Synchronized Gutter: Absolute line tracking & Jump-to-Error navigation",
        "6. Dual Themes: WCAG AA compliant Night Mode (Dark Theme)",
        "7. Verification: 136 passing automated regression & stress tests"
    ]
    for h in highlights:
        hp = htf.add_paragraph()
        hp.text = h
        hp.font.size = Pt(11)
        hp.font.color.rgb = RGBColor(226, 232, 240)
        hp.space_before = Pt(7)

    # ==========================================
    # SLIDE 2: Evolution from Review 2 to Review 3
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, BG_LIGHT)
    add_header(slide2, "Evolution: From Review 2 to Final Review 3", "Project Progression")

    # Card 1: Review 2 Baseline
    c1 = slide2.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0)
    )
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = CARD_BORDER
    tf_c1 = c1.text_frame
    tf_c1.word_wrap = True
    tf_c1.margin_left = tf_c1.margin_right = Inches(0.35)
    tf_c1.margin_top = Inches(0.35)

    p = tf_c1.paragraphs[0]
    p.text = "REVIEW 2 BASELINE (INITIAL STATE)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = WARNING_AMBER

    r2_items = [
        "• Single-Language Subset: Limited only to a small C subset (main, types, printf).",
        "• Small Program Capacity: Tested only on short snippets (~10 to 50 lines).",
        "• Static Pipeline Output: IR and assembly were generated, but code could NOT actually run.",
        "• Truncation & Buffer Freezing: Large pastes were truncated to ~96 lines by clipboard buffers.",
        "• Line Number Gutter Clamping: Line numbers desynchronized and stayed stuck at lines 1..27.",
        "• Single Light Theme: Incomplete dark mode with unreadable white-on-pale AST nodes.",
        "• Baseline Test Suite: 17 basic unit tests covering single-language endpoints."
    ]
    for item in r2_items:
        p = tf_c1.add_paragraph()
        p.text = item
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_NAVY
        p.space_before = Pt(10)

    # Arrow between cards
    arr = slide2.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW,
        Inches(6.45), Inches(3.8), Inches(0.4), Inches(0.4)
    )
    arr.fill.solid()
    arr.fill.fore_color.rgb = PRIMARY_BLUE
    arr.line.fill.background()

    # Card 2: Review 3 Upgrades
    c2 = slide2.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(6.9), Inches(1.8), Inches(5.6), Inches(5.0)
    )
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = SUCCESS_GREEN
    tf_c2 = c2.text_frame
    tf_c2.word_wrap = True
    tf_c2.margin_left = tf_c2.margin_right = Inches(0.35)
    tf_c2.margin_top = Inches(0.35)

    p = tf_c2.paragraphs[0]
    p.text = "REVIEW 3 DELIVERABLES (COMPLETED)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = SUCCESS_GREEN

    r3_items = [
        "• 6 Fully-Fledged Frontends: C, C++, Java, Python, JavaScript, and Toy Lang with isolated presets.",
        "• 5,000+ Line Scalability: Document processing up to 5,000 lines in <0.2s with 10k line limits.",
        "• Sandboxed Execution Engine: In-memory TAC VM (20k step budget) + safe Python runner.",
        "• Lossless Ingestion: Dedicated clipboard paste handler + multi-language source file upload.",
        "• Synchronized Viewport Gutter: Absolute tracking, dynamic digit widths, and Jump-to-Error sync.",
        "• WCAG AA Contrast Night Theme: Complete dark mode design tokens across all components.",
        "• 136 Automated Tests: 100% passing rate across 6 stress and regression test suites."
    ]
    for item in r3_items:
        p = tf_c2.add_paragraph()
        p.text = item
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_NAVY
        p.space_before = Pt(10)

    # ==========================================
    # SLIDE 3: System Architecture (Polyglot Pipeline)
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, BG_LIGHT)
    add_header(slide3, "Rev 3 Polyglot Compiler & Execution Architecture", "System Architecture")

    # Layer 1: Ingestion
    ingest_box = slide3.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.8), Inches(1.8), Inches(11.733), Inches(0.85)
    )
    ingest_box.fill.solid()
    ingest_box.fill.fore_color.rgb = RGBColor(238, 242, 255)
    ingest_box.line.color.rgb = PURPLE_ACCENT
    itf = ingest_box.text_frame
    itf.margin_top = Inches(0.12)
    p = itf.paragraphs[0]
    p.text = "1. INGESTION LAYER (Large-File & Polyglot Input)"
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = PURPLE_ACCENT
    p2 = itf.add_paragraph()
    p2.text = "Upload Source File (.c, .cpp, .java, .js, .py, .toy, .txt)  |  Lossless Clipboard Interceptor  |  Language-Isolated Example Presets"
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_NAVY

    # Layer 2: 6 Language Frontends
    front_box = slide3.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.8), Inches(2.8), Inches(11.733), Inches(1.1)
    )
    front_box.fill.solid()
    front_box.fill.fore_color.rgb = RGBColor(240, 249, 255)
    front_box.line.color.rgb = PRIMARY_BLUE
    ftf = front_box.text_frame
    ftf.margin_top = Inches(0.12)
    p = ftf.paragraphs[0]
    p.text = "2. PLUGGABLE FRONTEND LAYER (get_compiler(lang) Dispatcher)"
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = PRIMARY_BLUE
    p2 = ftf.add_paragraph()
    p2.text = "C Compiler (C99 Subset)  |  C++ Compiler (Classes & I/O)  |  Java Compiler (Main & Methods)\nPython Compiler (Recursive ast.parse)  |  JavaScript Compiler (ES6 Subset)  |  Toy Lang Compiler"
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_NAVY

    # Layer 3: Middle-End Pipeline
    mid_box = slide3.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.8), Inches(4.05), Inches(11.733), Inches(1.1)
    )
    mid_box.fill.solid()
    mid_box.fill.fore_color.rgb = RGBColor(240, 253, 244)
    mid_box.line.color.rgb = SUCCESS_GREEN
    mtf = mid_box.text_frame
    mtf.margin_top = Inches(0.12)
    p = mtf.paragraphs[0]
    p.text = "3. INTERMEDIATE ANALYSIS & CODE GENERATION LAYER"
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = SUCCESS_GREEN
    p2 = mtf.add_paragraph()
    p2.text = "Normalized Recursive AST  ->  Symbol Table & Scoping  ->  Linear TAC & Quadruples  ->  x86-64 & RISC Assembly Emission\nHigh-Speed O(N) Token Grouping  |  Stage-by-Stage Semantic Diagnostic Reporting  |  Jump-to-Error Offset Resolution"
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_NAVY

    # Layer 4: Dual Execution Engines
    exec_box = slide3.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.8), Inches(5.3), Inches(11.733), Inches(1.35)
    )
    exec_box.fill.solid()
    exec_box.fill.fore_color.rgb = CARD_BG
    exec_box.line.color.rgb = CARD_BORDER
    etf = exec_box.text_frame
    etf.margin_top = Inches(0.12)
    p = etf.paragraphs[0]
    p.text = "4. SAFE EXECUTION & PEDAGOGICAL PRESENTATION LAYER"
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_NAVY

    p2 = etf.add_paragraph()
    p2.text = "• TAC Virtual Machine (tac_vm.py): Sandboxed in-memory simulation with a 20,000 instruction budget (cross-platform, zero host compiler dependencies)."
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = TEXT_MUTED

    p3 = etf.add_paragraph()
    p3.text = "• Safe Python Native Runner (SafeExecutor): Subprocess isolation with 2.0s timeout guard and fallback to TAC VM."
    p3.font.size = Pt(10.5)
    p3.font.color.rgb = TEXT_MUTED

    p4 = etf.add_paragraph()
    p4.text = "• UI Features: Interactive Stepper, Stage Diff Mode, Token Pills, AST Visual Explorer, Execution Terminal, WCAG AA Dark Theme."
    p4.font.size = Pt(10.5)
    p4.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 4: Multi-Language Compiler Pipelines (6 Languages)
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, BG_LIGHT)
    add_header(slide4, "Six Supported Languages & Feature Matrix", "Pluggable Frontends")

    # Table of 6 Languages
    rows = 7
    cols = 5
    top = Inches(1.8)
    left = Inches(0.8)
    width = Inches(11.733)
    height = Inches(4.8)

    table_shape = slide4.shapes.add_table(rows, cols, left, top, width, height)
    tbl = table_shape.table

    tbl.columns[0].width = Inches(1.8)
    tbl.columns[1].width = Inches(3.0)
    tbl.columns[2].width = Inches(2.8)
    tbl.columns[3].width = Inches(2.2)
    tbl.columns[4].width = Inches(1.933)

    headers = ["Language", "Grammar & Construct Coverage", "AST & Semantic Validation", "IR & Target Assembly", "Execution Engine"]
    for c_idx, h in enumerate(headers):
        cell = tbl.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_NAVY
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = RGBColor(255, 255, 255)

    data = [
        ("C (C99 Subset)", "Prototypes, pointers, arrays, const/unsigned, printf", "Scope resolution, uninitialized vars, headers", "Three-Address Code, x86-64, RISC", "TAC VM (Sandboxed)"),
        ("C++", "Classes, public methods, constructors, cout/cin", "Member access, typo suggestions (coutt -> cout)", "Object TAC instructions, x86-64", "TAC VM (Sandboxed)"),
        ("Java", "Class definitions, public static void main, methods", "Type checking, typo diagnostics (println)", "Class-level TAC quadruples, x86-64", "TAC VM (Sandboxed)"),
        ("Python", "Classes, functions, nested loops, if/else, collections", "Full recursive ast.parse, undefined var check", "Pythonic TAC transformations, RISC", "Native Runner + TAC VM"),
        ("JavaScript", "Functions, let/const, binary ops, console.log", "Identifier declarations, typo hints (conosle)", "Linear TAC & x86-64 assembly", "TAC VM (Sandboxed)"),
        ("Toy Lang", "Minimal educational syntax: vars, expressions, print", "Strict variable declaration before use", "Pedagogical TAC quadruples", "TAC VM (Sandboxed)")
    ]

    for r_idx, row_data in enumerate(data, start=1):
        for c_idx, val in enumerate(row_data):
            cell = tbl.cell(r_idx, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG if r_idx % 2 == 1 else CARD_BG_ALT
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(10)
            p.font.bold = (c_idx == 0)
            p.font.color.rgb = PRIMARY_BLUE if c_idx == 0 else TEXT_NAVY

    # ==========================================
    # SLIDE 5: 5,000-Line Large Input Scalability & Stress Benchmarks
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, BG_LIGHT)
    add_header(slide5, "Large-Input Reliability: 5,000+ Lines Benchmarked", "Scalability & Performance")

    # Benchmark Table
    rows = 7
    cols = 5
    top = Inches(1.8)
    left = Inches(0.8)
    width = Inches(7.5)
    height = Inches(4.8)

    table_shape = slide5.shapes.add_table(rows, cols, left, top, width, height)
    tbl = table_shape.table

    tbl.columns[0].width = Inches(1.5)
    tbl.columns[1].width = Inches(1.3)
    tbl.columns[2].width = Inches(1.6)
    tbl.columns[3].width = Inches(1.6)
    tbl.columns[4].width = Inches(1.5)

    headers = ["Language", "Line Count", "Full Pipeline Status", "Output Produced", "Measured Time"]
    for c_idx, h in enumerate(headers):
        cell = tbl.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_NAVY
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = RGBColor(255, 255, 255)

    bench_data = [
        ("C (C99)", "5,002 lines", "Passed (All 6 Stages)", "Total: 3", "0.089 s  (89 ms)"),
        ("C++", "5,002 lines", "Passed (All 6 Stages)", "Result: 3", "0.073 s  (73 ms)"),
        ("Java", "5,001 lines", "Passed (All 6 Stages)", "Java: 3", "0.074 s  (74 ms)"),
        ("Python", "5,000 lines", "Passed (All 6 Stages)", "Python: 3", "0.190 s  (190 ms)"),
        ("JavaScript", "5,001 lines", "Passed (All 6 Stages)", "JS: 3", "0.070 s  (70 ms)"),
        ("Toy Lang", "5,000 lines", "Passed (All 6 Stages)", "3", "0.159 s  (159 ms)")
    ]

    for r_idx, row_data in enumerate(bench_data, start=1):
        for c_idx, val in enumerate(row_data):
            cell = tbl.cell(r_idx, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG if r_idx % 2 == 1 else CARD_BG_ALT
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(10.5)
            p.font.bold = (c_idx == 0 or c_idx == 4)
            if c_idx == 4:
                p.font.color.rgb = SUCCESS_GREEN
            elif c_idx == 0:
                p.font.color.rgb = PRIMARY_BLUE
            else:
                p.font.color.rgb = TEXT_NAVY

    # Right Side Optimization Summary Card
    opt_card = slide5.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(8.6), Inches(1.8), Inches(3.933), Inches(4.8)
    )
    opt_card.fill.solid()
    opt_card.fill.fore_color.rgb = CARD_BG
    opt_card.line.color.rgb = CARD_BORDER
    opt_tf = opt_card.text_frame
    opt_tf.word_wrap = True
    opt_tf.margin_left = opt_tf.margin_right = Inches(0.3)
    opt_tf.margin_top = Inches(0.3)

    p = opt_tf.paragraphs[0]
    p.text = "SCALING ARCHITECTURE"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE

    opts = [
        "• Linear Token Indexation: High-speed O(N) grouping of tokens by line number eliminates DOM quadratic freezing.",
        "• Memory-Bound AST Construction: Normalized single-pass tree builders parse 5,000 statements without stack overflow.",
        "• Step-Budgeted VM Execution: 20,000 step instruction cap strictly prevents infinite loops while executing complex logic.",
        "• 16 MB Payload Guard: Flask config['MAX_CONTENT_LENGTH'] configured to support up to 10,000-line source payloads.",
        "• High-Stress Accuracy: Deliberate syntax error at line 4,900 pinpointed with exact line 4900, col 24 diagnostic."
    ]
    for opt in opts:
        p = opt_tf.add_paragraph()
        p.text = opt
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_NAVY
        p.space_before = Pt(8)

    # ==========================================
    # SLIDE 6: Source Workspace Engineering & Line Number Synchronization
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, BG_LIGHT)
    add_header(slide6, "Source Workspace: Lossless Input & Synchronized Gutter", "Engineering Fixes")

    # Column 1: The Root Cause Bug
    c1 = slide6.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0)
    )
    c1.fill.solid()
    c1.fill.fore_color.rgb = CARD_BG
    c1.line.color.rgb = WARNING_AMBER
    tf1 = c1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_right = Inches(0.35)
    tf1.margin_top = Inches(0.3)

    p = tf1.paragraphs[0]
    p.text = "THE ROOT CAUSE (BEFORE FIX)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = WARNING_AMBER

    before_bullets = [
        "1. Clipboard Pipe Truncation: Shell / OS clipboard buffers clipped large pastes to exactly 4,096 bytes (~96 lines), causing premature EOF errors.",
        "2. Flexbox Gutter Sizing Bug: .line-numbers-gutter had no height restriction inside a flex container. 5,000 divs expanded the gutter to 100,727px.",
        "3. Zero Scroll Range Clamping: Because gutter clientHeight == scrollHeight (100,727px), maxScrollTop was 0. gutter.scrollTop was clamped to 0.",
        "4. Visual Symptom: The textarea scrolled to the bottom (line 5000), but the gutter remained frozen at lines 1..27.",
        "5. DOM Thrashing: Re-rendering 5,000 DOM divs on every keystroke froze the browser UI thread."
    ]
    for b in before_bullets:
        p = tf1.add_paragraph()
        p.text = b
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_NAVY
        p.space_before = Pt(8)

    # Column 2: Architectural Solution
    c2 = slide6.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(6.9), Inches(1.8), Inches(5.6), Inches(5.0)
    )
    c2.fill.solid()
    c2.fill.fore_color.rgb = CARD_BG
    c2.line.color.rgb = SUCCESS_GREEN
    tf2 = c2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_right = Inches(0.35)
    tf2.margin_top = Inches(0.3)

    p = tf2.paragraphs[0]
    p.text = "THE ARCHITECTURAL SOLUTION (IMPLEMENTED)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = SUCCESS_GREEN

    after_bullets = [
        "1. Lossless Paste Interceptor: Custom paste event listener reads e.clipboardData.getData('text') directly, guaranteeing 100% paste retention.",
        "2. Native File Upload Feature: Direct file ingestion for .c, .cpp, .java, .js, .py, .toy, .txt with automatic language detection.",
        "3. Absolute Viewport Tracking: Positioned gutter absolutely (top: 0, bottom: 0, left: 0), governing gutter height 100% by the visible textarea.",
        "4. Dynamic Gutter Width: Width scales dynamically from 46px (1-99 lines) to 68px (10,000+ lines) with proportional editor padding.",
        "5. Synchronized Error Navigation: Jump-to-Error sets synchronized scroll offsets for both editor and gutter to exact line targets (e.g. line 4,900)."
    ]
    for b in after_bullets:
        p = tf2.add_paragraph()
        p.text = b
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_NAVY
        p.space_before = Pt(8)

    # ==========================================
    # SLIDE 7: Night Theme & WCAG AA Contrast Compliance
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, BG_LIGHT)
    add_header(slide7, "UI/UX & Night Theme: WCAG AA Contrast Compliance", "User Experience")

    # Three feature cards
    card_width = Inches(3.64)
    gap = Inches(0.4)

    # Card 1: Design Tokens
    k1 = slide7.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.8), Inches(1.8), card_width, Inches(5.0)
    )
    k1.fill.solid()
    k1.fill.fore_color.rgb = CARD_BG
    k1.line.color.rgb = CARD_BORDER
    tf_k1 = k1.text_frame
    tf_k1.word_wrap = True
    tf_k1.margin_left = tf_k1.margin_right = Inches(0.25)
    tf_k1.margin_top = Inches(0.25)

    p = tf_k1.paragraphs[0]
    p.text = "1. Centralized Theme Tokens"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = PRIMARY_BLUE

    items_k1 = [
        "• Unified CSS variable architecture supporting instant light/dark toggle.",
        "• Backgrounds: #FFFFFF / #0E1422 with radial grid overlay.",
        "• Surfaces: Theme-aware surface elevation cards with subtle 1px border contrast.",
        "• Typography Tokens: Strict text-primary, text-secondary, and text-muted contrast ratios.",
        "• Interactive Toggle: One-click moon/sun icon toggle in header with localStorage persistence."
    ]
    for it in items_k1:
        p = tf_k1.add_paragraph()
        p.text = it
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_NAVY
        p.space_before = Pt(8)

    # Card 2: AST Node Contrast
    k2 = slide7.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.8) + card_width + gap, Inches(1.8), card_width, Inches(5.0)
    )
    k2.fill.solid()
    k2.fill.fore_color.rgb = CARD_BG
    k2.line.color.rgb = CARD_BORDER
    tf_k2 = k2.text_frame
    tf_k2.word_wrap = True
    tf_k2.margin_left = tf_k2.margin_right = Inches(0.25)
    tf_k2.margin_top = Inches(0.25)

    p = tf_k2.paragraphs[0]
    p.text = "2. AST Node Contrast Fix"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = SUCCESS_GREEN

    items_k2 = [
        "• Problem in Rev 2: Pale node backgrounds washed out white text in Dark Theme.",
        "• Node Palettes Redesigned: Program, FunctionDef, VarDecl, BinOp, Return nodes given dedicated dark/light tokens.",
        "• Dynamic Label Colors: Node headers and attributes compute dynamic high-contrast foregrounds.",
        "• WCAG AA Compliance: All 7 AST node classes verified >= 4.5:1 text-to-background contrast ratio.",
        "• Tested in Automated Suite: Verified via test_wcag_aa_contrast_for_all_ast_nodes."
    ]
    for it in items_k2:
        p = tf_k2.add_paragraph()
        p.text = it
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_NAVY
        p.space_before = Pt(8)

    # Card 3: Guided Explainer Panels
    k3 = slide7.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.8) + (card_width + gap) * 2, Inches(1.8), card_width, Inches(5.0)
    )
    k3.fill.solid()
    k3.fill.fore_color.rgb = CARD_BG
    k3.line.color.rgb = CARD_BORDER
    tf_k3 = k3.text_frame
    tf_k3.word_wrap = True
    tf_k3.margin_left = tf_k3.margin_right = Inches(0.25)
    tf_k3.margin_top = Inches(0.25)

    p = tf_k3.paragraphs[0]
    p.text = "3. Explainer & Diff Usability"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = PURPLE_ACCENT

    items_k3 = [
        "• Key Takeaway Panels: Restyled with dark-surface backgrounds and luminous amber/teal text.",
        "• Code Viewers: Monospace code viewer components adapt with dark-mode syntax highlight styling.",
        "• Stage Diff Mode: Side-by-side AST and TAC diff tables styled with legible green/red additions/deletions.",
        "• Diagnostic Chips: Error and warning chips dynamically styled for maximum readability in both modes.",
        "• Line Gutter Indicators: Error lines highlighted with red border markers and amber warnings."
    ]
    for it in items_k3:
        p = tf_k3.add_paragraph()
        p.text = it
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_NAVY
        p.space_before = Pt(8)

    # ==========================================
    # SLIDE 8: Demonstration Scenarios (Review 3 Live Workflows)
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8, BG_LIGHT)
    add_header(slide8, "Live Demonstration: End-to-End Key Scenarios", "Live Verification")

    demo_steps = [
        ("01", "Polyglot Compilation & Execution",
         "Select Python/Java/C++ preset or upload file. Execute full 6-stage pipeline (Scanner -> Parser -> Semantics -> TAC -> Assembly -> VM Execution) and view terminal output."),
        ("02", "5,000-Line Stress Test Ingestion",
         "Load a 5,000-line C/Python benchmark via 'Upload Source File' or lossless paste. Watch the editor and gutter synchronize up to line 5,000 and compile in <0.2s."),
        ("03", "Large-Line Error Diagnostic & Jump",
         "Introduce a deliberate syntax error at line 4,900. View precise diagnostic chip and click 'Jump to Error' to synchronously navigate both editor and gutter to line 4,900."),
        ("04", "Multi-Language Stage Diff & Night Theme",
         "Toggle between Light and Night themes. Compare Program A vs Program B in Stage Diff Mode across AST and Three-Address Code with full WCAG AA contrast.")
    ]

    for idx, (num, title, desc) in enumerate(demo_steps):
        top_offset = Inches(1.8 + idx * 1.25)
        
        # Number badge
        nb = slide8.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.8), top_offset, Inches(0.9), Inches(1.05)
        )
        nb.fill.solid()
        nb.fill.fore_color.rgb = DARK_NAVY
        nb.line.color.rgb = PRIMARY_BLUE
        ntf = nb.text_frame
        ntf.margin_top = Inches(0.2)
        np = ntf.paragraphs[0]
        np.text = num
        np.font.size = Pt(22)
        np.font.bold = True
        np.font.color.rgb = PRIMARY_BLUE
        np.alignment = PP_ALIGN.CENTER

        # Content Card
        cc = slide8.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(1.85), top_offset, Inches(10.68), Inches(1.05)
        )
        cc.fill.solid()
        cc.fill.fore_color.rgb = CARD_BG
        cc.line.color.rgb = CARD_BORDER
        ctf = cc.text_frame
        ctf.word_wrap = True
        ctf.margin_left = Inches(0.25)
        ctf.margin_top = Inches(0.12)

        cp = ctf.paragraphs[0]
        cp.text = title
        cp.font.size = Pt(13)
        cp.font.bold = True
        cp.font.color.rgb = TEXT_NAVY

        cp2 = ctf.add_paragraph()
        cp2.text = desc
        cp2.font.size = Pt(10.5)
        cp2.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 9: Testing & Verification: 136 Tests Across 6 Suites
    # ==========================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9, BG_LIGHT)
    add_header(slide9, "Comprehensive Verification: 136 Automated Tests", "Testing & Reliability")

    # Big KPI Stat Block on Left
    kpi_card = slide9.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0)
    )
    kpi_card.fill.solid()
    kpi_card.fill.fore_color.rgb = DARK_NAVY
    kpi_card.line.color.rgb = PRIMARY_BLUE
    ktf = kpi_card.text_frame
    ktf.word_wrap = True
    ktf.margin_top = Inches(0.4)

    kp1 = ktf.paragraphs[0]
    kp1.text = "136 / 136"
    kp1.font.size = Pt(44)
    kp1.font.bold = True
    kp1.font.color.rgb = PRIMARY_BLUE
    kp1.alignment = PP_ALIGN.CENTER

    kp2 = ktf.add_paragraph()
    kp2.text = "TESTS PASSING (100%)"
    kp2.font.size = Pt(13)
    kp2.font.bold = True
    kp2.font.color.rgb = RGBColor(255, 255, 255)
    kp2.alignment = PP_ALIGN.CENTER

    kp3 = ktf.add_paragraph()
    kp3.text = "\nProgression:\n• Review 1: 8 Tests\n• Review 2: 17 Tests\n• Review 3: 136 Tests\n  (+700% Test Expansion)"
    kp3.font.size = Pt(11)
    kp3.font.color.rgb = RGBColor(203, 213, 225)
    kp3.alignment = PP_ALIGN.CENTER

    kp4 = ktf.add_paragraph()
    kp4.text = "\nExecution Time: 13.30s\nZero Failing Tests"
    kp4.font.size = Pt(12)
    kp4.font.bold = True
    kp4.font.color.rgb = SUCCESS_GREEN
    kp4.alignment = PP_ALIGN.CENTER

    # Test Suites Breakdown Table on Right
    rows = 7
    cols = 4
    top = Inches(1.8)
    left = Inches(4.7)
    width = Inches(7.833)
    height = Inches(5.0)

    t_table = slide9.shapes.add_table(rows, cols, left, top, width, height)
    tt = t_table.table

    tt.columns[0].width = Inches(2.8)
    tt.columns[1].width = Inches(0.9)
    tt.columns[2].width = Inches(3.133)
    tt.columns[3].width = Inches(1.0)

    s_headers = ["Test Suite File", "Tests", "Scope & Verifications", "Status"]
    for c_idx, h in enumerate(s_headers):
        cell = tt.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_NAVY
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(10.5)
        p.font.color.rgb = RGBColor(255, 255, 255)

    suites_data = [
        ("test_pipeline.py", "43", "Core 6-stage pipeline, Lexer, Parser, AST, TAC, ASM", "PASSED"),
        ("test_review3_upgrades.py", "42", "6 languages, TAC VM, Night theme, WCAG AA contrast", "PASSED"),
        ("test_final_review_stress.py", "28", "3k & 4k line benchmarks, high-line diagnostics", "PASSED"),
        ("test_diff_mode.py", "11", "Side-by-side AST & TAC diff visual comparisons", "PASSED"),
        ("test_paste_and_large_input_fix.py", "8", "5,000-line C stress test, parser prototypes & types", "PASSED"),
        ("test_gutter_line_numbers.py", "4", "Browser DOM viewport gutter tracking (Edge CDP)", "PASSED")
    ]

    for r_idx, row in enumerate(suites_data, start=1):
        for c_idx, val in enumerate(row):
            cell = tt.cell(r_idx, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG if r_idx % 2 == 1 else CARD_BG_ALT
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(9.5)
            p.font.bold = (c_idx == 0 or c_idx == 3)
            if c_idx == 3:
                p.font.color.rgb = SUCCESS_GREEN
            elif c_idx == 0:
                p.font.color.rgb = PRIMARY_BLUE
            else:
                p.font.color.rgb = TEXT_NAVY

    # ==========================================
    # SLIDE 10: Conclusion & Summary
    # ==========================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10, DARK_NAVY)

    # Title
    t_box = slide10.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(1.0))
    tf10 = t_box.text_frame
    p = tf10.paragraphs[0]
    p.text = "Conclusion: A Production-Ready Educational Compiler"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    bar = slide10.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.8), Inches(1.7), Inches(1.5), Inches(0.06)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = PRIMARY_BLUE
    bar.line.color.rgb = PRIMARY_BLUE

    # Summary Grid (4 Cards)
    cw = Inches(5.6)
    ch = Inches(2.2)

    quads = [
        ("Multi-Language Polyglot Support",
         "Successfully eliminated single-language restrictions. Added 6 genuine language frontends (C, C++, Java, Python, JavaScript, Toy Lang) with dedicated presets, ASTs, and semantic diagnostics.",
         Inches(0.8), Inches(2.1), PRIMARY_BLUE),
        ("5,000+ Line Scalability & Speed",
         "Benchmarked and verified reliable processing of 5,000+ lines in all 6 languages under 200ms. Eliminated quadratic DOM thrashing and supported safe payloads up to 10,000 lines.",
         Inches(6.9), Inches(2.1), SUCCESS_GREEN),
        ("In-Memory TAC Virtual Machine",
         "Introduced sandboxed virtual execution for Three-Address Code with instruction budgeting. Students and faculty can now observe actual program execution without installing native compilers.",
         Inches(0.8), Inches(4.6), PURPLE_ACCENT),
        ("Engineering Reliability & Night Theme",
         "Lossless paste handler, multi-file upload, synchronized absolute viewport gutter, WCAG AA contrast compliance in Light/Dark themes, and 136 automated regression tests.",
         Inches(6.9), Inches(4.6), WARNING_AMBER)
    ]

    for title, desc, left, top, accent in quads:
        card = slide10.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            left, top, cw, ch
        )
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(30, 41, 59)
        card.line.color.rgb = accent
        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = Inches(0.3)
        ctf.margin_top = Inches(0.2)

        p = ctf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = accent

        p2 = ctf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = RGBColor(226, 232, 240)
        p2.space_before = Pt(6)

    # Save presentation
    output_path = "Review_3_Compiler_Pipeline_Visualizer.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path} ({os.path.getsize(output_path)} bytes)")

if __name__ == "__main__":
    create_deck()
